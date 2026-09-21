"""Generic Excel (.xlsx) parser supporting multi-sheet workbooks and setup templates."""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import openpyxl

from daily_expenses_importer.parsers.base import BaseParser, NormalizedTransaction, TransactionType

_DATE_HEADERS = {"date", "fecha", "fecha y hora", "fecha/hora", "datetime", "date & time", "date and time", "transaction date", "fecha transaccion", "fecha transacción"}
_DESC_HEADERS = {"description", "descripción", "descripcion", "concepto", "detalle", "merchant", "payee", "comercio", "notas", "nota"}
_AMOUNT_HEADERS = {"amount", "monto", "importe", "valor"}
_DEBIT_HEADERS = {"debit", "debito", "débito", "gasto", "egreso", "cargos", "expense", "expenses"}
_CREDIT_HEADERS = {"credit", "credito", "crédito", "ingreso", "abono", "abonos", "income", "incomes"}
_REF_HEADERS = {"reference", "referencia", "ref", "id", "secuencia", "número", "numero", "nro operacion"}
_CAT_HEADERS = {"category", "categoría", "categoria", "rubro", "clasificación", "clasificacion"}
_TYPE_HEADERS = {"type", "tipo", "tipo movimiento", "tipo de movimiento", "tipo_movimiento"}
_ACCOUNT_HEADERS = {"account", "cuenta", "target account", "target_account", "cuenta destino", "cuenta_destino"}
_SOURCE_ACC_HEADERS = {"source account", "source_account", "cuenta origen", "cuenta_origen"}
_TIME_HEADERS = {"time", "hora", "horario"}

_DATA_SHEET_KEYWORDS = {"movimiento", "movement", "transaccion", "transaction", "dato", "gasto", "registro"}
_SETUP_SHEET_KEYWORDS = {"instruccion", "instruction", "configura", "setup", "cuenta", "ayuda", "info"}


class GenericExcelParser(BaseParser):
    """Parses standard Excel (.xlsx) workbooks with automatic sheet and header detection."""

    @property
    def name(self) -> str:
        return "generic_excel"

    def supports(self, file_path: Path, sample_bytes: bytes) -> bool:
        return file_path.suffix.lower() in {".xlsx", ".xlsm"}

    def parse(self, file_path: Path) -> list[NormalizedTransaction]:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception:
            return []

        try:
            sheet = self._select_data_sheet(wb)
            if sheet is None:
                return []

            rows = list(sheet.iter_rows(values_only=True))
        finally:
            wb.close()

        if not rows:
            return []

        # Find header row
        header_idx, field_map = self._find_header_row(rows)
        if header_idx is None or ("date" not in field_map or ("amount" not in field_map and "debit" not in field_map)):
            return []

        transactions = []
        for idx, row in enumerate(rows[header_idx + 1:], 1):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            tx = self._parse_row(row, field_map, fallback_index=idx)
            if tx is not None:
                transactions.append(tx)

        return transactions

    def _select_data_sheet(self, wb: openpyxl.Workbook):
        """Select the sheet containing movement data."""
        # 1. Look for sheet with data keywords
        for name in wb.sheetnames:
            clean = name.strip().lower()
            if any(k in clean for k in _DATA_SHEET_KEYWORDS):
                return wb[name]

        # 2. If first sheet is a setup sheet, choose the second sheet
        if len(wb.sheetnames) > 1:
            first_clean = wb.sheetnames[0].strip().lower()
            if any(k in first_clean for k in _SETUP_SHEET_KEYWORDS):
                return wb[wb.sheetnames[1]]

        # 3. Fallback to active sheet
        return wb.active

    def _find_header_row(self, rows: list[tuple]) -> tuple[int | None, dict[str, int]]:
        """Scan top rows to find header column indices."""
        for idx, row in enumerate(rows[:10]):
            mapped = {}
            for col_idx, val in enumerate(row):
                if val is None:
                    continue
                clean = str(val).strip().lower().replace("_", " ")
                if clean in _DATE_HEADERS and "date" not in mapped:
                    mapped["date"] = col_idx
                elif clean in _DESC_HEADERS and "desc" not in mapped:
                    mapped["desc"] = col_idx
                elif clean in _AMOUNT_HEADERS and "amount" not in mapped:
                    mapped["amount"] = col_idx
                elif clean in _DEBIT_HEADERS and "debit" not in mapped:
                    mapped["debit"] = col_idx
                elif clean in _CREDIT_HEADERS and "credit" not in mapped:
                    mapped["credit"] = col_idx
                elif clean in _REF_HEADERS and "ref" not in mapped:
                    mapped["ref"] = col_idx
                elif clean in _CAT_HEADERS and "category" not in mapped:
                    mapped["category"] = col_idx
                elif clean in _TYPE_HEADERS and "type" not in mapped:
                    mapped["type"] = col_idx
                elif clean in _ACCOUNT_HEADERS and "account" not in mapped:
                    mapped["account"] = col_idx
                elif clean in _SOURCE_ACC_HEADERS and "source_account" not in mapped:
                    mapped["source_account"] = col_idx
                elif clean in _TIME_HEADERS and "time" not in mapped:
                    mapped["time"] = col_idx

            if "date" in mapped and ("amount" in mapped or "debit" in mapped):
                return idx, mapped

        return None, {}

    def _parse_row(self, row: tuple, field_map: dict[str, int], fallback_index: int) -> NormalizedTransaction | None:
        raw_date_val = row[field_map["date"]] if field_map.get("date") is not None and field_map["date"] < len(row) else None
        parsed_date, extracted_time = self._parse_datetime(raw_date_val)
        if not parsed_date:
            return None

        # Description
        raw_desc = ""
        if field_map.get("desc") is not None and field_map["desc"] < len(row):
            raw_desc = str(row[field_map["desc"]] or "").strip()

        # Reference
        ref = ""
        if field_map.get("ref") is not None and field_map["ref"] < len(row):
            ref = str(row[field_map["ref"]] or "").strip()
        if not ref:
            ref = f"XLSX-{parsed_date.strftime('%Y%m%d')}-{fallback_index}"

        # Amount & Transaction Type
        amount = Decimal("0")
        tx_type = TransactionType.EXPENSE

        if "debit" in field_map or "credit" in field_map:
            debit_raw = row[field_map["debit"]] if field_map.get("debit") is not None and field_map["debit"] < len(row) else None
            credit_raw = row[field_map["credit"]] if field_map.get("credit") is not None and field_map["credit"] < len(row) else None
            debit_val = self._to_decimal(debit_raw)
            credit_val = self._to_decimal(credit_raw)
            if debit_val > Decimal("0"):
                amount = debit_val
                tx_type = TransactionType.EXPENSE
            elif credit_val > Decimal("0"):
                amount = credit_val
                tx_type = TransactionType.INCOME
            else:
                return None
        elif "amount" in field_map:
            amt_raw = row[field_map["amount"]] if field_map["amount"] < len(row) else None
            amt_val = self._to_decimal(amt_raw)
            if amt_val < Decimal("0") or (amt_raw and "-" in str(amt_raw)):
                amount = abs(amt_val)
                tx_type = TransactionType.EXPENSE
            elif amt_val > Decimal("0"):
                amount = amt_val
                tx_type = TransactionType.INCOME
            else:
                return None
        else:
            return None

        # Type column override
        if field_map.get("type") is not None and field_map["type"] < len(row):
            raw_type = str(row[field_map["type"]] or "").strip().lower()
            if "transfer" in raw_type or "transf" in raw_type:
                tx_type = TransactionType.TRANSFER
            elif "inc" in raw_type or "ingr" in raw_type or "abono" in raw_type:
                tx_type = TransactionType.INCOME
            elif "exp" in raw_type or "gast" in raw_type or "egr" in raw_type:
                tx_type = TransactionType.EXPENSE

        # Category
        category = None
        if field_map.get("category") is not None and field_map["category"] < len(row):
            cat_val = str(row[field_map["category"]] or "").strip()
            category = cat_val if cat_val else None

        # Account
        account_raw = ""
        if field_map.get("account") is not None and field_map["account"] < len(row):
            account_raw = str(row[field_map["account"]] or "").strip()

        source_acc = None
        if field_map.get("source_account") is not None and field_map["source_account"] < len(row):
            src_val = str(row[field_map["source_account"]] or "").strip()
            source_acc = src_val if src_val else None

        target_acc = account_raw or None

        # Time column or extracted time
        time_str = extracted_time
        if field_map.get("time") is not None and field_map["time"] < len(row):
            raw_time_col = row[field_map["time"]]
            if isinstance(raw_time_col, time):
                time_str = raw_time_col.strftime("%H:%M:%S")
            elif raw_time_col:
                time_str = str(raw_time_col).strip()

        # Handle transfer accounts format: "Source -> Target"
        if account_raw and "->" in account_raw:
            tx_type = TransactionType.TRANSFER
            parts = [p.strip() for p in account_raw.split("->", 1)]
            source_acc = parts[0]
            target_acc = parts[1]

        return NormalizedTransaction(
            reference=ref,
            date=parsed_date,
            time=time_str,
            amount=amount,
            currency="USD",
            description=raw_desc,
            raw_description=raw_desc,
            tx_type=tx_type,
            category=category,
            source_account=source_acc,
            target_account=target_acc,
        )

    def _parse_datetime(self, val) -> tuple[date | None, str]:
        """Extract (date, time_string) from Excel cell value."""
        if val is None:
            return None, ""

        if isinstance(val, datetime):
            time_str = val.strftime("%H:%M:%S") if (val.hour or val.minute or val.second) else ""
            return val.date(), time_str
        elif isinstance(val, date):
            return val, ""

        text = str(val).strip()
        if not text:
            return None, ""

        # Check for combined date and time strings
        time_str = ""
        formats = [
            ("%Y-%m-%d %H:%M:%S", True),
            ("%Y-%m-%d %H:%M", True),
            ("%d/%m/%Y %H:%M:%S", True),
            ("%d/%m/%Y %H:%M", True),
            ("%Y/%m/%d %H:%M:%S", True),
            ("%Y/%m/%d %H:%M", True),
            ("%Y-%m-%d", False),
            ("%d/%m/%Y", False),
            ("%m/%d/%Y", False),
            ("%d-%m-%Y", False),
            ("%Y/%m/%d", False),
            ("%d.%m.%Y", False),
        ]

        for fmt, has_time in formats:
            try:
                dt = datetime.strptime(text, fmt)
                if has_time:
                    time_str = dt.strftime("%H:%M:%S")
                return dt.date(), time_str
            except ValueError:
                continue

        # Split space fallback
        parts = text.split(" ")
        if len(parts) >= 2:
            time_part = parts[1]
            try:
                d = datetime.strptime(parts[0], "%Y-%m-%d").date()
                return d, time_part
            except Exception:
                pass

        return None, ""

    def _to_decimal(self, val) -> Decimal:
        if val is None:
            return Decimal("0")
        if isinstance(val, (int, float, Decimal)):
            return Decimal(str(val))
        text = str(val).replace("$", "").replace("€", "").replace("USD", "").replace(" ", "").strip()
        is_negative = False
        if text.startswith("(") and text.endswith(")"):
            text = text[1:-1].strip()
            is_negative = True
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            dec = Decimal(text)
            return -dec if is_negative else dec
        except Exception:
            return Decimal("0")
