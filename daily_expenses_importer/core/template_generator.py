"""Template generator for Daily Expenses 4 Importer.

Generates:
1. `template.xlsx`: Accounting Ledger multi-sheet workbook with separate Expense & Income columns,
   dynamic dropdown validations switching categories automatically, and synced accounts.
2. `template.csv`: Matching Accounting Ledger plain text CSV (date,description,expense,income,category,account).
"""

from __future__ import annotations

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

from daily_expenses_importer.core.metadata_fetcher import load_cached_metadata

DEFAULT_ACCOUNTS = [
    "Checking Account",
    "Credit Card",
    "Cash",
    "Savings",
    "Payroll",
]

DEFAULT_EXPENSE_CATEGORIES = [
    "Bills",
    "Clothing",
    "Dining",
    "Education",
    "Entertainment",
    "Fuel",
    "Groceries",
    "Healthcare",
    "Home",
    "Insurance",
    "Kids",
    "Merchandise",
    "Other",
    "Personal",
    "Pets",
    "Subscriptions",
    "Tips",
    "Toll",
    "Transportation",
    "Travel",
    "Utilities",
    "Vehicle",
]

DEFAULT_INCOME_CATEGORIES = [
    "Bonus",
    "Freelance",
    "Gift",
    "Interest",
    "Investment",
    "Loan Repayment",
    "Other Income",
    "Refund",
    "Rental",
    "Salary",
    "Sales",
]

# Combined default categories list for backwards compatibility
DEFAULT_CATEGORIES = sorted(list(set(DEFAULT_EXPENSE_CATEGORIES + DEFAULT_INCOME_CATEGORIES)))


def _resolve_template_lists(
    accounts: list[str] | None = None,
    expense_categories: list[str] | None = None,
    income_categories: list[str] | None = None,
) -> tuple[list[str], list[str], list[str]]:
    """Resolve accounts and categories from arguments, cached metadata, or defaults."""
    cached = load_cached_metadata() or {}

    res_accounts = accounts or cached.get("accounts") or DEFAULT_ACCOUNTS
    res_expense_cats = expense_categories or cached.get("expense_categories") or DEFAULT_EXPENSE_CATEGORIES
    res_income_cats = income_categories or cached.get("income_categories") or DEFAULT_INCOME_CATEGORIES

    return res_accounts, res_expense_cats, res_income_cats


def _find_best_match(candidates: list[str], preferred: list[str], fallback: str) -> str:
    """Find the best matching category or account name based on preferred keywords."""
    for pref in preferred:
        for c in candidates:
            if pref.lower() in c.lower():
                return c
    return candidates[0] if candidates else fallback


def generate_excel_template(
    dest_path: Path = Path("template.xlsx"),
    accounts: list[str] | None = None,
    expense_categories: list[str] | None = None,
    income_categories: list[str] | None = None,
) -> Path:
    """Generate a rich, Accounting Ledger template.xlsx with dynamic dropdown validations."""
    acc_list, exp_cats, inc_cats = _resolve_template_lists(accounts, expense_categories, income_categories)

    wb = openpyxl.Workbook()

    # Style definitions
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # -------------------------------------------------------------
    # SHEET 1: Setup (Instructions, Accounts, Expense & Income Categories)
    # -------------------------------------------------------------
    ws_setup = wb.active
    ws_setup.title = "Setup"

    ws_setup["B2"] = "🏦 Daily Expenses 4 Importer — Setup & Categories"
    ws_setup["B2"].font = title_font

    ws_setup["B3"] = "Instructions:"
    ws_setup["B3"].font = Font(name="Calibri", size=11, bold=True)
    instructions = [
        "1. Column B lists your active Accounts (synced from your Daily Expenses 4 app).",
        "2. Column D lists your Expense Categories; Column F lists your Income Categories.",
        "3. Go to the 'Movements' sheet and enter transactions in Accounting Ledger format (Expense vs Income).",
        "4. When you enter an amount in 'Income', the Category dropdown automatically shows Income categories!",
        "5. For transfers, enter the amount in 'Expense' and use 'Source -> Target' in the 'Account' column.",
        "6. In 'Date & Time', use YYYY-MM-DD HH:MM:SS format (e.g. 2026-09-15 14:30:00).",
        "7. To import, run: daily-expenses-importer -i template.xlsx",
    ]
    for i, inst in enumerate(instructions, 4):
        ws_setup[f"B{i}"] = inst
        ws_setup[f"B{i}"].font = subtitle_font

    # Tables Headers
    ws_setup["B12"] = "Accounts"
    ws_setup["B12"].font = section_font
    ws_setup["B12"].fill = section_fill
    ws_setup["B12"].alignment = Alignment(horizontal="center")

    ws_setup["D12"] = "Expense Categories"
    ws_setup["D12"].font = section_font
    ws_setup["D12"].fill = section_fill
    ws_setup["D12"].alignment = Alignment(horizontal="center")

    ws_setup["F12"] = "Income Categories"
    ws_setup["F12"].font = section_font
    ws_setup["F12"].fill = section_fill
    ws_setup["F12"].alignment = Alignment(horizontal="center")

    # Populate accounts
    for idx, acc in enumerate(acc_list, 13):
        cell = ws_setup[f"B{idx}"]
        cell.value = acc
        cell.border = thin_border

    # Populate expense categories
    for idx, cat in enumerate(exp_cats, 13):
        cell = ws_setup[f"D{idx}"]
        cell.value = cat
        cell.border = thin_border

    # Populate income categories
    for idx, cat in enumerate(inc_cats, 13):
        cell = ws_setup[f"F{idx}"]
        cell.value = cat
        cell.border = thin_border

    ws_setup.column_dimensions["A"].width = 4
    ws_setup.column_dimensions["B"].width = 30
    ws_setup.column_dimensions["C"].width = 4
    ws_setup.column_dimensions["D"].width = 30
    ws_setup.column_dimensions["E"].width = 4
    ws_setup.column_dimensions["F"].width = 30

    # -------------------------------------------------------------
    # SHEET 2: Movements (Accounting Ledger Format)
    # -------------------------------------------------------------
    ws_data = wb.create_sheet(title="Movements")

    headers = [
        "Date & Time",
        "Description",
        "Expense",
        "Income",
        "Category",
        "Account",
    ]

    for col_num, h_text in enumerate(headers, 1):
        cell = ws_data.cell(row=1, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Default samples adapted to real or default values
    sample_acc = _find_best_match(acc_list, ["checking", "corriente", "personal", "main"], "Checking Account")
    sample_acc2 = _find_best_match([a for a in acc_list if a != sample_acc], ["credit", "tarjeta", "tdc", "card"], "Credit Card")
    sample_exp_cat = _find_best_match(exp_cats, ["groceries", "food", "comida", "mercado"], "Groceries")
    sample_exp_cat2 = _find_best_match(exp_cats, ["health", "pharmacy", "salud", "farmacia", "medicina"], "Healthcare")
    sample_exp_dining = _find_best_match(exp_cats, ["dining", "restaurant", "restaurante"], sample_exp_cat)
    sample_inc_cat = _find_best_match(inc_cats, ["salary", "sueldo", "salario", "nomina"], "Salary")

    sample_rows = [
        ("2026-09-15 14:30:00", "Supermarket Groceries", 45.50, None, sample_exp_cat, sample_acc),
        ("2026-09-16 09:15:00", "Monthly Salary Deposit", None, 1500.00, sample_inc_cat, sample_acc),
        ("2026-09-17 18:20:00", "Credit Card Payment", 120.00, None, None, f"{sample_acc} -> {sample_acc2}"),
        ("2026-09-18 11:45:00", "Pharmacy & Medicine", 12.30, None, sample_exp_cat2, sample_acc),
        ("2026-09-19 20:30:00", "Restaurant Dinner", 34.00, None, sample_exp_dining, sample_acc2),
    ]

    for row_idx, row_data in enumerate(sample_rows, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in (3, 4) and val is not None:
                cell.number_format = "$#,##0.00"

    # Data Validations (Dropdowns)
    # 1. Dynamic Category Validation (Income vs Expense)
    last_exp_row = 12 + len(exp_cats)
    last_inc_row = 12 + len(inc_cats)
    dv_cat_formula = f"=IF(ISNUMBER($D2), Setup!$F$13:$F${last_inc_row}, Setup!$D$13:$D${last_exp_row})"
    dv_cat = DataValidation(type="list", formula1=dv_cat_formula, allow_blank=True)
    dv_cat.error = "Select a category from the list or configure it in the Setup sheet"
    dv_cat.errorTitle = "Invalid Category"
    ws_data.add_data_validation(dv_cat)
    dv_cat.add("E2:E500")

    # 2. Account Validation referencing Sheet 1
    last_acc_row = 12 + len(acc_list)
    dv_acc = DataValidation(type="list", formula1=f"=Setup!$B$13:$B${last_acc_row}", allow_blank=True)
    dv_acc.error = "Select an account from the list or configure it in the Setup sheet"
    dv_acc.errorTitle = "Invalid Account"
    ws_data.add_data_validation(dv_acc)
    dv_acc.add("F2:F500")

    # Column widths
    ws_data.column_dimensions["A"].width = 24
    ws_data.column_dimensions["B"].width = 32
    ws_data.column_dimensions["C"].width = 16
    ws_data.column_dimensions["D"].width = 16
    ws_data.column_dimensions["E"].width = 25
    ws_data.column_dimensions["F"].width = 35

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest_path)
    wb.close()
    return dest_path


def generate_csv_template(
    dest_path: Path = Path("template.csv"),
    accounts: list[str] | None = None,
    expense_categories: list[str] | None = None,
    income_categories: list[str] | None = None,
) -> Path:
    """Generate a standard Accounting Ledger template.csv file."""
    acc_list, exp_cats, inc_cats = _resolve_template_lists(accounts, expense_categories, income_categories)

    sample_acc = _find_best_match(acc_list, ["checking", "corriente", "personal", "main"], "Checking Account")
    sample_acc2 = _find_best_match([a for a in acc_list if a != sample_acc], ["credit", "tarjeta", "tdc", "card"], "Credit Card")
    sample_exp_cat = _find_best_match(exp_cats, ["groceries", "food", "comida", "mercado"], "Groceries")
    sample_exp_cat2 = _find_best_match(exp_cats, ["health", "pharmacy", "salud", "farmacia", "medicina"], "Healthcare")
    sample_exp_dining = _find_best_match(exp_cats, ["dining", "restaurant", "restaurante"], sample_exp_cat)
    sample_inc_cat = _find_best_match(inc_cats, ["salary", "sueldo", "salario", "nomina"], "Salary")

    content = f"""date,description,expense,income,category,account
2026-09-15 14:30:00,Supermarket Groceries,45.50,,{sample_exp_cat},{sample_acc}
2026-09-16 09:15:00,Monthly Salary Deposit,,1500.00,{sample_inc_cat},{sample_acc}
2026-09-17 18:20:00,Credit Card Payment,120.00,,,{sample_acc} -> {sample_acc2}
2026-09-18 11:45:00,Pharmacy & Medicine,12.30,,{sample_exp_cat2},{sample_acc}
2026-09-19 20:30:00,Restaurant Dinner,34.00,,{sample_exp_dining},{sample_acc2}
"""
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    dest_path.write_text(content, encoding="utf-8")
    return dest_path
